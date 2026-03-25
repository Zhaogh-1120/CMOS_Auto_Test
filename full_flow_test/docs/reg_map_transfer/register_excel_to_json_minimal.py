#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
将“寄存器迁移模板_AutoTest_最小输入版.xlsx”转换为工程所需的：
- register_fields.json
- registers.json

适配工作簿页签：
- 字段输入模板
- 可选_寄存器汇总
- 项目配置（可选，只做信息读取，不影响核心输出）

设计目标：
1. 允许 Excel 公式未缓存时也能正常转换（脚本自行计算关键字段）
2. 尽量容错：支持 bit_spec 为 "7:3" / "2" / "<7:3>" / "<2>"
3. 输出格式与当前工程 register_fields.json / registers.json 保持兼容
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

FIELD_SHEET = "字段输入模板"
REGISTER_SHEET = "可选_寄存器汇总"
PROJECT_SHEET = "项目配置"


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    return _clean_text(value) == ""


def _to_int(value: Any, field_name: str = "") -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        raise ValueError(f"字段 {field_name} 不是整数: {value}")

    text = str(value).strip()
    if text == "":
        return None

    # 兼容十六进制输入
    if text.lower().startswith("0x"):
        return int(text, 16)

    # 兼容诸如 32.0
    if re.fullmatch(r"[-+]?\d+\.0+", text):
        return int(float(text))

    if re.fullmatch(r"[-+]?\d+", text):
        return int(text)

    raise ValueError(f"字段 {field_name} 不是可识别的整数: {value}")


def _parse_bit_spec(bit_spec: Any) -> tuple[int, int, str, int, int]:
    text = _clean_text(bit_spec)
    if not text:
        raise ValueError("bit_spec 为空")

    text = text.replace(" ", "")
    text = text.strip("<>")

    m = re.fullmatch(r"(\d+):(\d+)", text)
    if m:
        msb = int(m.group(1))
        lsb = int(m.group(2))
    else:
        m = re.fullmatch(r"(\d+)", text)
        if not m:
            raise ValueError(f"bit_spec 格式错误: {bit_spec}")
        msb = lsb = int(m.group(1))

    if msb < lsb:
        raise ValueError(f"bit_spec 非法，msb < lsb: {bit_spec}")

    width = msb - lsb + 1
    mask = ((1 << width) - 1) << lsb
    bits = f"<{msb}>" if msb == lsb else f"<{msb}:{lsb}>"
    return msb, lsb, bits, width, mask


def _hex8(value: int) -> str:
    if value < 0:
        raise ValueError(f"地址不能为负数: {value}")
    return f"0x{value:02X}"


def _read_header_map(ws, header_row: int) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        name = _clean_text(ws.cell(header_row, col).value)
        if name:
            header_map[name] = col
    return header_map


def _cell(ws, row: int, header_map: dict[str, int], header_name: str) -> Any:
    col = header_map.get(header_name)
    if not col:
        return None
    return ws.cell(row, col).value


def parse_register_sheet(wb) -> OrderedDict[str, dict[str, Any]]:
    ws = wb[REGISTER_SHEET]
    header_map = _read_header_map(ws, 4)

    required = ["register_name", "address_decimal"]
    for key in required:
        if key not in header_map:
            raise KeyError(f"页签“{REGISTER_SHEET}”缺少列: {key}")

    registers: OrderedDict[str, dict[str, Any]] = OrderedDict()
    errors: list[str] = []

    for row in range(5, ws.max_row + 1):
        register_name = _clean_text(_cell(ws, row, header_map, "register_name"))
        address_raw = _cell(ws, row, header_map, "address_decimal")
        desc = _clean_text(_cell(ws, row, header_map, "description"))
        notes = _clean_text(_cell(ws, row, header_map, "notes"))

        if not register_name and _is_blank(address_raw) and not desc and not notes:
            continue
        if not register_name:
            errors.append(f"{REGISTER_SHEET}!A{row}: register_name 为空")
            continue

        try:
            address = _to_int(address_raw, f"{REGISTER_SHEET}:{register_name}:address_decimal")
        except Exception as exc:
            errors.append(f"{REGISTER_SHEET}!{row}: {exc}")
            continue

        if address is None:
            errors.append(f"{REGISTER_SHEET}!{row}: register {register_name} 缺少 address_decimal")
            continue

        if register_name in registers and registers[register_name]["addr"] != address:
            errors.append(
                f"{REGISTER_SHEET}!{row}: register {register_name} 地址冲突，"
                f"已有 {registers[register_name]['addr']}，当前 {address}"
            )
            continue

        registers[register_name] = {
            "addr": address,
            "desc": desc,
        }

    if errors:
        raise ValueError("\n".join(errors))

    return registers


def parse_field_sheet(wb, registers: OrderedDict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    ws = wb[FIELD_SHEET]
    header_map = _read_header_map(ws, 3)

    required = ["register_name", "field_base_name", "bit_spec", "default_decimal"]
    for key in required:
        if key not in header_map:
            raise KeyError(f"页签“{FIELD_SHEET}”缺少列: {key}")

    fields: list[dict[str, Any]] = []
    field_name_seen: set[str] = set()
    alias_seen: dict[str, str] = {}
    errors: list[str] = []

    # 只把前 8 列看作“是否整行为空”的判断区域，和模板逻辑一致
    input_headers = [
        "register_name",
        "field_base_name",
        "alias",
        "bit_spec",
        "default_decimal",
        "access",
        "description",
        "notes",
    ]

    for row in range(4, ws.max_row + 1):
        input_values = [_cell(ws, row, header_map, h) for h in input_headers]
        if all(_is_blank(v) for v in input_values):
            continue

        register_name = _clean_text(_cell(ws, row, header_map, "register_name"))
        field_base_name = _clean_text(_cell(ws, row, header_map, "field_base_name"))
        alias = _clean_text(_cell(ws, row, header_map, "alias")) or field_base_name
        bit_spec = _cell(ws, row, header_map, "bit_spec")
        default_raw = _cell(ws, row, header_map, "default_decimal")

        access = _clean_text(_cell(ws, row, header_map, "access"))
        description = _clean_text(_cell(ws, row, header_map, "description"))
        notes = _clean_text(_cell(ws, row, header_map, "notes"))

        row_label = f"{FIELD_SHEET}!{row}"

        if not register_name:
            errors.append(f"{row_label}: register_name 为空")
            continue
        if register_name not in registers:
            errors.append(f"{row_label}: register_name={register_name} 未在“{REGISTER_SHEET}”中找到")
            continue
        if not field_base_name:
            errors.append(f"{row_label}: field_base_name 为空")
            continue

        try:
            msb, lsb, bits, width, mask = _parse_bit_spec(bit_spec)
            default_decimal = _to_int(default_raw, f"{row_label}:default_decimal")
        except Exception as exc:
            errors.append(f"{row_label}: {exc}")
            continue

        if default_decimal is None:
            errors.append(f"{row_label}: default_decimal 为空")
            continue
        if default_decimal < 0:
            errors.append(f"{row_label}: default_decimal 不能为负数")
            continue

        max_decimal = (1 << width) - 1
        if default_decimal > max_decimal:
            errors.append(
                f"{row_label}: default_decimal={default_decimal} 超出字段范围 0~{max_decimal}"
            )
            continue

        address = registers[register_name]["addr"]
        field_full_name = f"{field_base_name}{bits}"

        if field_full_name in field_name_seen:
            errors.append(f"{row_label}: 字段全名重复 {field_full_name}")
            continue
        field_name_seen.add(field_full_name)

        # alias 允许重复，但记录冲突信息，便于使用者避坑
        alias_key = alias.upper()
        if alias_key in alias_seen and alias_seen[alias_key] != field_full_name:
            print(
                f"[WARN] alias 重复：{alias} -> {alias_seen[alias_key]} / {field_full_name}，"
                f"运行时按当前工程逻辑可能造成 alias 歧义。",
                file=sys.stderr,
            )
        else:
            alias_seen[alias_key] = field_full_name

        field_item = {
            "name": field_full_name,
            "alias": alias,
            "address": address,
            "address_hex": _hex8(address),
            "bits": bits,
            "msb": msb,
            "lsb": lsb,
            "width": width,
            "mask": mask,
            "default_decimal": default_decimal,
            "max_decimal": max_decimal,
        }

        # 这些信息不写入最终 json，但保留给调试报告/后续扩展
        field_item["_meta"] = {
            "register_name": register_name,
            "access": access,
            "description": description,
            "notes": notes,
        }
        fields.append(field_item)

    if errors:
        raise ValueError("\n".join(errors))

    return fields


def build_registers_json(registers: OrderedDict[str, dict[str, Any]]) -> OrderedDict[str, dict[str, Any]]:
    out: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for reg_name, reg_info in registers.items():
        out[reg_name] = {
            "addr": reg_info["addr"],
            "desc": reg_info.get("desc", ""),
        }
    return out


def build_register_fields_json(fields: list[dict[str, Any]]) -> list[dict[str, Any]]:
    # 保持与当前工程 FieldDef 兼容；去掉 _meta
    out = []
    for item in fields:
        out.append({k: v for k, v in item.items() if not k.startswith("_")})
    return out


def compute_register_defaults(fields: list[dict[str, Any]]) -> OrderedDict[int, int]:
    defaults: OrderedDict[int, int] = OrderedDict()
    for item in fields:
        addr = item["address"]
        defaults.setdefault(addr, 0)
        defaults[addr] = (defaults[addr] & ~item["mask"]) | ((item["default_decimal"] << item["lsb"]) & item["mask"])
    return defaults


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def write_report(path: Path, registers_json: OrderedDict[str, dict[str, Any]], fields: list[dict[str, Any]]) -> None:
    defaults = compute_register_defaults(fields)
    lines = []
    lines.append("# 寄存器迁移转换报告")
    lines.append("")
    lines.append(f"- 寄存器数量: {len(registers_json)}")
    lines.append(f"- 字段数量: {len(fields)}")
    lines.append("")
    lines.append("## 寄存器默认值汇总")
    lines.append("")
    lines.append("| register_name | addr_dec | addr_hex | default_dec | default_hex | desc |")
    lines.append("|---|---:|---|---:|---|---|")
    for reg_name, reg_info in registers_json.items():
        addr = reg_info["addr"]
        default_dec = defaults.get(addr, 0)
        desc = reg_info.get("desc", "") or ""
        lines.append(
            f"| {reg_name} | {addr} | {_hex8(addr)} | {default_dec} | {_hex8(default_dec)} | {desc.replace('|','/')} |"
        )
    lines.append("")
    lines.append("## 字段列表")
    lines.append("")
    lines.append("| name | alias | address | bits | default_decimal |")
    lines.append("|---|---|---:|---|---:|")
    for item in fields:
        lines.append(
            f"| {item['name']} | {item['alias']} | {item['address']} | {item['bits']} | {item['default_decimal']} |"
        )
    path.write_text("\n".join(lines), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="将最小输入版寄存器 Excel 模板转换为 register_fields.json 和 registers.json"
    )
    parser.add_argument("excel", help="Excel 模板路径")
    parser.add_argument(
        "--out-dir",
        default=".",
        help="输出目录，默认当前目录。会写出 register_fields.json / registers.json / conversion_report.md",
    )
    parser.add_argument(
        "--fields-name",
        default="register_fields.json",
        help="输出的字段 JSON 文件名，默认 register_fields.json",
    )
    parser.add_argument(
        "--registers-name",
        default="registers.json",
        help="输出的寄存器 JSON 文件名，默认 registers.json",
    )
    parser.add_argument(
        "--report-name",
        default="conversion_report.md",
        help="输出报告文件名，默认 conversion_report.md",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    excel_path = Path(args.excel).expanduser().resolve()
    out_dir = Path(args.out_dir).expanduser().resolve()

    if not excel_path.exists():
        print(f"[ERROR] Excel 文件不存在: {excel_path}", file=sys.stderr)
        return 2

    try:
        wb = load_workbook(excel_path, data_only=False)
    except Exception as exc:
        print(f"[ERROR] 打开 Excel 失败: {exc}", file=sys.stderr)
        return 2

    for sheet_name in [FIELD_SHEET, REGISTER_SHEET]:
        if sheet_name not in wb.sheetnames:
            print(f"[ERROR] 缺少必要页签: {sheet_name}", file=sys.stderr)
            return 2

    try:
        registers = parse_register_sheet(wb)
        fields = parse_field_sheet(wb, registers)
        registers_json = build_registers_json(registers)
        register_fields_json = build_register_fields_json(fields)
    except Exception as exc:
        print("[ERROR] 转换失败：", file=sys.stderr)
        print(str(exc), file=sys.stderr)
        return 1

    write_json(out_dir / args.fields_name, register_fields_json)
    write_json(out_dir / args.registers_name, registers_json)
    write_report(out_dir / args.report_name, registers_json, fields)

    print(f"[OK] 已生成: {out_dir / args.fields_name}")
    print(f"[OK] 已生成: {out_dir / args.registers_name}")
    print(f"[OK] 已生成: {out_dir / args.report_name}")
    print(f"[INFO] 寄存器数量: {len(registers_json)}")
    print(f"[INFO] 字段数量: {len(register_fields_json)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
